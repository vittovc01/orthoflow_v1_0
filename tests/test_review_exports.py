import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook

from review_exports import review_csv, review_xlsx


def test_exports_preserve_current_fields_multiplicity_and_text_identifiers():
    rows = [dict(codice='04.200.030S', lotto='00123', scadenza='2035-01-01',
                 descrizione='Vite; titanio', quantita=2, gtin='technical', patient='not for export'),
            dict(codice='04.200.030S', lotto='00123', scadenza='2035-01-01',
                 descrizione='Seconda etichetta', quantita=1),
            dict(codice='EXCLUDED', quantita=1, escludi_riga=True)]
    csv_rows = list(csv.DictReader(StringIO(review_csv(rows).decode('utf-8-sig')), delimiter=';'))
    assert list(csv_rows[0]) == ['REF','Lotto','Scadenza','Descrizione','Quantità']
    assert len(csv_rows) == 2 and csv_rows[0]['Lotto'] == '00123'
    assert csv_rows[0]['Descrizione'] == 'Vite; titanio'
    book = load_workbook(BytesIO(review_xlsx(rows)))
    sheet = book['Materiali']
    assert sheet.max_row == 3 and sheet['B2'].value == '00123'
    assert sheet['A2'].value == '04.200.030S' and sheet['E2'].value == 2
    assert sheet['E3'].value == 1
    assert 'technical' not in str(list(sheet.values)) and 'not for export' not in str(list(sheet.values))


def test_exports_are_literal_and_retain_incomplete_fields_for_review():
    rows = [dict(codice='TEST', lotto='LOT', scadenza='', descrizione='=1+1', quantita='1.5')]
    book = load_workbook(BytesIO(review_xlsx(rows)))
    assert book.active['D2'].value == '=1+1' and book.active['D2'].data_type == 's'
    assert book.active['E2'].value == '1.5'
    assert "'=1+1" in review_csv(rows).decode('utf-8-sig')
