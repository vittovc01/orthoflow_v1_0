import io
import json
import pytest
from fastapi.testclient import TestClient
from mobile.app import create_app, COOKIE
from mobile.service import code_key
from tests.fixtures.mobile_backend import Backend


@pytest.fixture
def setup(tmp_path,monkeypatch):
    monkeypatch.setenv('MOBILE_PUBLIC_ORIGIN','https://testserver')
    monkeypatch.delenv('MOBILE_LOCAL_DEV',raising=False)
    backend=Backend()
    app=create_app(lambda:backend,tmp_path/'sessions.sqlite')
    client=TestClient(app,base_url='https://testserver')
    client.headers['Origin']='https://testserver'
    return client,backend,app


def login(client,name='mario'):
    return client.post('/api/login',json={'username':name,'password':'test-password'})


def payload(**changes):
    data=dict(request_id='test-request-id-123456',customer_code='9010013',warehouse='MAG1',line='TRAUMA',procedure_date='2026-10-03',items=[dict(code='1504-00-126',lot='B',quantity=2)])
    data.update(changes)
    return data


def test_cookie_is_secure_and_contains_no_identity(setup):
    client,backend,app=setup
    response=login(client)
    assert response.status_code==200
    assert 'Secure' in response.headers['set-cookie'] and 'HttpOnly' in response.headers['set-cookie']
    assert 'SameSite=strict' in response.headers['set-cookie']
    assert not any(k in response.json() for k in ['password_hash','password_salt','id'])
    token=client.cookies[COOKIE]
    assert 'mario' not in token
    assert app.state.sessions.get(token)[0]==1


def test_user_cannot_see_or_modify_another_couriers_mission(setup):
    client,backend,_=setup
    login(client)
    assert [m['id'] for m in client.get('/api/missions').json()['missions']]==[101]
    for method,path,body in [('GET','/api/missions/102',None),('POST','/api/missions/102/stamp',{'latitude':41,'longitude':14,'accuracy':5}),('POST','/api/missions/102/complete',{})]:
        assert client.request(method,path,json=body).status_code==404
    assert backend.calls==[]


def test_revocation_and_changed_permissions_apply_to_live_session(setup):
    client,backend,_=setup
    login(client)
    backend.tables['utenti_app'][0]['permessi']=[]
    assert client.get('/api/missions').status_code==403
    backend.tables['utenti_app'][0]['attivo']=False
    assert client.get('/api/me').status_code==401


def test_password_change_revokes_old_session(setup):
    client,backend,_=setup
    login(client)
    backend.tables['utenti_app'][0]['password_hash']='changed'
    assert client.get('/api/me').status_code==401


def test_cross_origin_write_and_forged_roles_are_rejected(setup):
    client,backend,_=setup
    client.headers['Origin']='https://attacker.invalid'
    assert login(client).status_code==403
    client.headers['Origin']='https://testserver'
    assert client.post('/api/login',json={'username':'mario','password':'test-password','ruolo':'Admin'}).status_code==422


def test_login_rate_limit(setup):
    client,_,_=setup
    for _ in range(15): assert login(client).status_code==200
    assert login(client).status_code==429


def test_stamp_validates_coordinates_and_ignores_forged_identity(setup):
    client,backend,_=setup
    login(client)
    assert client.post('/api/missions/101/stamp',json={'latitude':400,'longitude':14,'accuracy':5}).status_code==422
    assert client.post('/api/missions/101/stamp',json={'latitude':41,'longitude':14,'accuracy':5,'user_id':2}).status_code==422
    assert client.post('/api/missions/101/stamp',json={'latitude':41,'longitude':14,'accuracy':5}).status_code==200
    assert backend.calls[-1][1]['p_user_id']==1


def test_agent_customers_prices_and_preserved_hyphen_codes(setup):
    client,backend,_=setup
    login(client,'agente')
    assert [c['codice_cliente'] for c in client.get('/api/scarico/options').json()['customers']]==['9010013']
    response=client.post('/api/scarico/preview',json=payload())
    assert response.status_code==200
    result=response.json()
    assert result['items'][0]['codice']=='1504-00-126'
    assert result['items'][0]['prezzo']==145 and result['total']==290
    assert result['items'][0]['prezzo_source']=='MANUALE_MEMORIZZATO'
    assert code_key('1504-00-126')==code_key('150400126')
    assert client.post('/api/scarico/save',json=payload(customer_code='other')).status_code==403


def test_agent_cannot_forge_prices_or_omit_lot_confirmation(setup):
    client,backend,_=setup
    login(client,'agente')
    assert client.post('/api/scarico/preview',json=payload(items=[dict(code='x',lot='B',quantity=1,manual_price=42)])).status_code==403
    assert client.post('/api/scarico/preview',json=payload(items=[dict(code='x',quantity=1)])).status_code==422
    assert client.post('/api/scarico/preview',json=payload(items=[dict(code='x',quantity=1,lotless=True)])).status_code==200
    assert client.post('/api/scarico/preview',json=payload(items=[dict(code='x',quantity=1,lotless=True,manufacturer='Other Company')])).status_code==422
    assert client.post('/api/scarico/preview',json=payload(items=[dict(code='x',quantity=1,lotless=True,manufacturer='Other Company',verified=True)])).status_code==200


def test_export_is_real_excel_and_formula_is_text(setup):
    from openpyxl import load_workbook
    client,_,_=setup
    login(client,'agente')
    response=client.post('/api/scarico/export',json={'items':[dict(code='=1+1',lot='B',quantity=2)]})
    assert response.status_code==200
    sheet=load_workbook(io.BytesIO(response.content)).active
    assert sheet['A2'].value=='=1+1' and sheet['A2'].data_type=='s'


def test_spoofed_image_never_uploaded(setup):
    client,backend,_=setup
    login(client)
    response=client.post('/api/missions/101/photos',files={'files':('fake.jpg',b'<script>secret</script>','image/jpeg')})
    assert response.status_code==422
    assert backend.storage.uploads==[]


def test_api_never_cached_and_worker_caches_only_public_shell(setup):
    client,_,_=setup
    login(client)
    assert client.get('/api/me').headers['cache-control']=='no-store'
    worker=client.get('/sw.js').text
    assert '!SHELL.includes(url.pathname)' in worker
    assert '/api/' not in json.loads(client.get('/static/manifest.webmanifest').text)['start_url']


def test_public_health_does_not_report_ready_without_server_config(tmp_path,monkeypatch):
    for name in ('SUPABASE_URL','SUPABASE_SERVICE_KEY','MOBILE_PUBLIC_ORIGIN'):
        monkeypatch.delenv(name,raising=False)
    client=TestClient(create_app(session_path=tmp_path/'unconfigured.sqlite'))
    response=client.get('/health')
    assert response.status_code==503 and response.json()['status']=='not_configured'


def test_logout_invalidates_previous_cookie(setup):
    client,backend,app=setup
    login(client)
    old=client.cookies[COOKIE]
    assert client.post('/api/logout',json={}).status_code==200
    client.cookies.set(COOKIE,old,domain='testserver.local',path='/')
    assert app.state.sessions.get(old) is None
    assert client.get('/api/me').status_code==401


def test_health_detects_backend_failure_without_leaking_details(setup,monkeypatch):
    client,backend,_=setup
    def fail(*args):raise RuntimeError('private credential or clinical data')
    monkeypatch.setattr(backend,'table',fail)
    response=client.get('/health')
    assert response.status_code==503 and response.json()['status']=='backend_unavailable'
    assert 'private' not in response.text


def test_concurrent_courier_sessions_keep_missions_separate(setup):
    from concurrent.futures import ThreadPoolExecutor
    _,backend,app=setup
    def inspect(name):
        with TestClient(app,base_url='https://testserver') as client:
            client.headers['Origin']='https://testserver'
            assert login(client,name).status_code==200
            return [m['id'] for m in client.get('/api/missions').json()['missions']]
    with ThreadPoolExecutor(max_workers=2) as pool:
        results=list(pool.map(inspect,['mario','luigi']))
    assert results==[[101],[102]]
